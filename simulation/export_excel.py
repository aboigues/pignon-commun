# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1.5"]
# ///
"""Exporte la simulation économique au format Excel : simulation/simulation-economique.xlsx.

Le classeur contient un modèle en formules (hypothèses modifiables, cotisation, projection du
fonds sur 10 ans) et les résultats Monte-Carlo de simulation.py, lus dans resultats/resultats.json.

Usage : uv run simulation/export_excel.py
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

ICI = Path(__file__).parent
POLICE = "Arial"
BLEU = "0000FF"
VERT = "008000"
F_TITRE = Font(name=POLICE, size=14, bold=True)
F_SOUS = Font(name=POLICE, size=11, bold=True)
F_TEXTE = Font(name=POLICE, size=10)
F_NOTE = Font(name=POLICE, size=9, italic=True, color="555555")
F_ENTETE = Font(name=POLICE, size=10, bold=True, color="FFFFFF")
F_SAISIE = Font(name=POLICE, size=10, color=BLEU)
F_LIEN = Font(name=POLICE, size=10, color=VERT)
F_TOTAL = Font(name=POLICE, size=10, bold=True)
FOND_ENTETE = PatternFill("solid", fgColor="2F5D50")
FOND_CLE = PatternFill("solid", fgColor="FFFF00")
FOND_TOTAL = PatternFill("solid", fgColor="E8F0EC")
FOND_SECTION = PatternFill("solid", fgColor="D9E6DF")
TRAIT = Side(style="thin", color="B7C9C0")
BORDURE = Border(bottom=TRAIT)
ENVELOPPE = Alignment(wrap_text=True, vertical="top")

EUR = '#,##0 "€";-#,##0 "€";"-"'
EUR2 = '#,##0.00 "€";-#,##0.00 "€";"-"'
EUR3 = '0.000 "€";-0.000 "€";"-"'
PCT = '0.0%;-0.0%;"-"'
NB = '#,##0;-#,##0;"-"'
DEC = '0.00;-0.00;"-"'

# (catégorie, nom, libellé, valeur, unité, format, clé, source ou commentaire)
HYPOTHESES = [
    ("Parc", "n_leger", "Vélos à usage léger (vélo classique)", 1200, "vélos", NB, False, "Page, bloc « Passer à l'échelle »"),
    ("Parc", "n_intensif", "Vélos à usage intensif (vélo-cargo électrique)", 800, "vélos", NB, False, "Page, bloc « Passer à l'échelle »"),
    ("Parc", "n_ateliers", "Ateliers du réseau", 40, "ateliers", NB, False, "Page, bloc « Passer à l'échelle »"),
    ("Usage", "km_leger", "Kilométrage annuel, usage léger", 2800, "km/an", NB, False, "Ministère de la Transition écologique et FUB : 6 à 13 km domicile-travail × 200 jours"),
    ("Usage", "km_intensif", "Kilométrage annuel, usage intensif", 5500, "km/an", NB, False, "Observatoire des Boîtes à Vélo 2024 : 24 km par jour de tournée"),
    ("Usage", "churn", "Locataires qui partent chaque année", 0.20, "part", PCT, True, "Hypothèse, à mesurer auprès de Véligo et Swapfiets"),
    ("Usage", "remplissage", "Vélos vacants reloués dans l'année", 0.75, "part", PCT, True, "Hypothèse, à mesurer auprès de Véligo et Swapfiets"),
    ("Usage", "impayes", "Impayés", 0.02, "part des recettes", PCT, False, "Hypothèse"),
    ("Usage", "fraude_part", "Usagers qui sous-déclarent leurs km", 0.05, "part", PCT, False, "Hypothèse"),
    ("Usage", "fraude_sous_declaration", "Kilomètres non déclarés par un fraudeur", 0.30, "part", PCT, False, "Hypothèse"),
    ("Vélos", "cout_leger", "Coût de fabrication labellisé, vélo classique", 440, "€", EUR, False, "Page, bloc « D'où vient la rente d'usage » (illustratif)"),
    ("Vélos", "cout_intensif", "Coût de fabrication labellisé, vélo-cargo électrique", 1170, "€", EUR, True, "Page (illustratif). Prix publics des cargos français : 3 500 à 7 000 € ; paramètre le plus influent"),
    ("Vélos", "vie_leger", "Durée de vie garantie, vélo classique", 10, "ans", NB, False, "Page"),
    ("Vélos", "vie_intensif", "Durée de vie garantie, vélo-cargo", 8, "ans", NB, False, "Page"),
    ("Vélos", "taux_financement", "Taux de financement de la rente par le fabricant", 0.0, "par an", PCT, False, "Page : 0 %. Scénario S11 : 5 %"),
    ("Réparation", "h_1000km_leger", "Heures d'atelier par 1 000 km, usage léger", 3.3 / 2.8, "h", DEC, False, "Page : 3,3 h pour 2 800 km"),
    ("Réparation", "h_1000km_intensif", "Heures d'atelier par 1 000 km, usage intensif", 6.4 / 5.5, "h", DEC, False, "Page : 6,4 h pour 5 500 km"),
    ("Réparation", "h_facteur", "Multiplicateur des heures (erreur d'estimation)", 1.0, "×", DEC, True, "1 = hypothèse de la page. Scénario S06 : 1,4"),
    ("Réparation", "taux_horaire", "Taux horaire moyen (N1, N2, N3)", 47, "€/h", EUR, False, "Référentiel réparateur : 42, 52 et 62 €/h"),
    ("Réparation", "vieillissement", "Heures supplémentaires par année d'âge du vélo", 0.04, "par an", PCT, False, "Hypothèse"),
    ("Réparation", "forfait_suivi", "Forfait de suivi par vélo loué", 120, "€/an", EUR, False, "Référentiel réparateur : 10 €/mois"),
    ("Réparation", "forfait_vacant", "Forfait de suivi d'un vélo vacant (part du forfait)", 0.30, "part", PCT, False, "Hypothèse"),
    ("Réparation", "score_moyen", "Score de durabilité moyen du réseau", 75, "sur 100", NB, False, "Page (illustratif)"),
    ("Réparation", "prime_rep", "Prime réparateur : part du socle à score 100", 0.25, "part", PCT, False, "Référentiel réparateur"),
    ("Réparation", "prime_fab", "Prime fabricant : part de la rente à score 75", 0.16, "part", PCT, False, "Calée sur la page (7 € et 25 €)"),
    ("Pièces", "pieces_km_leger", "Pièces d'usure, usage léger", 0.025, "€/km", EUR3, False, "Hypothèse : environ 70 €/an (chaîne, pneus, plaquettes)"),
    ("Pièces", "pieces_km_intensif", "Pièces d'usure, vélo-cargo électrique", 0.045, "€/km", EUR3, False, "Hypothèse : environ 250 €/an"),
    ("Pièces", "pieces_facteur", "Multiplicateur de la consommation de pièces", 1.0, "×", DEC, True, "Scénario S07 : 2"),
    ("Pièces", "batterie_cout", "Batterie de remplacement", 650, "€", EUR, False, "500 à 1 000 € selon les sources ; environ 1 000 cycles"),
    ("Pièces", "batterie_remplacements", "Remplacements de batterie par vie de vélo-cargo", 1.05, "nombre", DEC, False, "Somme des probabilités annuelles de simulation.py"),
    ("Pièces", "batterie_facteur", "Multiplicateur de la fréquence de remplacement", 1.0, "×", DEC, False, "Scénario S08 : 2"),
    ("Sinistres", "casse_leger", "Casse prématurée par an, usage léger", 0.015, "part des vélos loués", PCT, False, "Page"),
    ("Sinistres", "casse_intensif", "Casse prématurée par an, usage intensif", 0.03, "part des vélos loués", PCT, False, "Page"),
    ("Sinistres", "casse_cout", "Coût moyen d'une casse prématurée", 400, "€", EUR, False, "Page"),
    ("Sinistres", "casse_facteur", "Multiplicateur de la casse", 1.0, "×", DEC, False, "Scénarios S04 et S05"),
    ("Sinistres", "vol_leger", "Vol par an, vélo classique", 0.02, "part du parc", PCT, False, "FUB : 350 000 à 580 000 vols par an en France"),
    ("Sinistres", "vol_intensif", "Vol par an, vélo-cargo électrique", 0.04, "part du parc", PCT, False, "Jusqu'à 6 % des vélos électriques selon les estimations"),
    ("Sinistres", "assurance_taux", "Prime d'assurance vol du parc", 0.05, "part de la valeur", PCT, False, "Hypothèse, à demander à une mutuelle"),
    ("Sinistres", "assurance_franchise", "Franchise d'assurance par vol", 0.10, "part de la valeur", PCT, False, "Hypothèse"),
    ("Sinistres", "franchise_client_leger", "Franchise refacturée au locataire, vélo classique", 150, "€", EUR, False, "Hypothèse"),
    ("Sinistres", "franchise_client_intensif", "Franchise refacturée au locataire, vélo-cargo", 300, "€", EUR, False, "Hypothèse"),
    ("Sinistres", "recouvrement_franchise", "Part des franchises effectivement recouvrée", 0.60, "part", PCT, False, "Hypothèse"),
    ("Parc", "capteur_cout", "Capteur kilométrique d'un vélo mécanique", 25, "€", EUR, False, "Page, bloc 08 : 15 à 30 € (capteur Bluetooth scellé)"),
    ("Parc", "capteur_renouvellement", "Capteurs remplacés chaque année", 0.10, "part", PCT, False, "Hypothèse"),
    ("Parc", "recond_leger", "Remise en état à la restitution, vélo classique", 60, "€", EUR, False, "Hypothèse"),
    ("Parc", "recond_intensif", "Remise en état à la restitution, vélo-cargo", 120, "€", EUR, False, "Hypothèse"),
    ("Gestion", "gestion_taux", "Chargement de gestion", 0.06, "part des coûts", PCT, False, "Page"),
    ("Gestion", "gestion_fixe", "Coûts de structure fixes en plus", 0, "€/an", EUR, False, "Scénario S22 : 180 000 €"),
    ("Gestion", "provision", "Provision de sécurité", 0.08, "part des coûts", PCT, False, "Page"),
    ("Projection", "inflation", "Inflation des coûts", 0.02, "par an", PCT, False, "Hypothèse"),
    ("Projection", "revision", "Révision annuelle du tarif (1 = oui, 0 = non)", 1, "0 ou 1", NB, True, "Recommandation de la simulation : obligatoire"),
    ("Projection", "cap_hausse", "Hausse maximale du tarif par an", 0.10, "par an", PCT, False, "Hypothèse de gouvernance"),
    ("Projection", "cap_baisse", "Baisse maximale du tarif par an", 0.10, "par an", PCT, False, "Hypothèse de gouvernance"),
    ("Projection", "capital_initial", "Réserve de démarrage", 0, "€", EUR, False, "Recommandation : environ 205 € par vélo"),
    ("Page initiale", "page_leger", "Cotisation annuelle affichée avant correction, usage léger", 431, "€/an", EUR, False, "Page avant le 2026-09-22"),
    ("Page initiale", "page_intensif", "Cotisation annuelle affichée avant correction, usage intensif", 765, "€/an", EUR, False, "Page avant le 2026-09-22"),
]


def entetes(ws, ligne, titres, largeurs=None):
    for i, t in enumerate(titres, start=1):
        c = ws.cell(ligne, i, t)
        c.font, c.fill = F_ENTETE, FOND_ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if largeurs:
        for i, w in enumerate(largeurs, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def cellule(ws, ref, valeur, police=F_TEXTE, fmt=None, fond=None):
    c = ws[ref]
    c.value = valeur
    c.font = police
    if fmt:
        c.number_format = fmt
    if fond:
        c.fill = fond
    return c


def feuille_lisez_moi(wb, res):
    ws = wb.active
    ws.title = "Lisez-moi"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    lignes = [
        (F_TITRE, "Pignon commun — simulation économique du fonds commun en location"),
        (F_TEXTE, "Classeur généré par simulation/export_excel.py à partir de simulation/simulation.py. Document de référence : docs/simulation-economique.md."),
        (None, ""),
        (F_SOUS, "Contenu"),
        (F_TEXTE, "Hypothèses — tous les paramètres du modèle, modifiables. Chaque valeur porte un nom (colonne F) utilisé dans les formules."),
        (F_TEXTE, "Cotisation — coût annuel attendu par vélo et tarif de location, en formules : il se recalcule quand une hypothèse change."),
        (F_TEXTE, "Fonds 10 ans — projection du fonds en valeur attendue (sans aléa), avec inflation et révision annuelle du tarif."),
        (F_TEXTE, f"Scénarios — résultats Monte-Carlo des {len(res['scenarios'])} scénarios critiques ({res['trajectoires_simulees']:,} trajectoires de 10 ans, graine {res['graine']}).".replace(",", " ")),
        (F_TEXTE, "Sensibilité — classement des 22 paramètres selon leur influence sur la marge, et carte de la probabilité de ruine."),
        (F_TEXTE, "Vacance prudente — effet d'un tarif calculé en supposant moins de vélos loués que prévu, face à plusieurs réalités."),
        (None, ""),
        (F_SOUS, "Légende"),
        (Font(name=POLICE, size=10, color=BLEU), "Texte bleu : valeur saisie, à modifier pour tester une hypothèse (feuille Hypothèses)."),
        (F_TEXTE, "Texte noir : formule calculée dans la feuille."),
        (Font(name=POLICE, size=10, color=VERT), "Texte vert : formule qui reprend une valeur d'une autre feuille."),
        (F_TEXTE, "Fond jaune : hypothèse clé, parmi les plus influentes selon l'analyse de sensibilité."),
        (None, ""),
        (F_SOUS, "Ce que ce classeur fait et ne fait pas"),
        (F_TEXTE, "Les feuilles Cotisation et Fonds 10 ans calculent des valeurs attendues : elles montrent l'effet moyen d'une hypothèse, pas le risque."),
        (F_TEXTE, "Le risque (probabilité de ruine, capital requis, tarif d'équilibre) vient de la simulation Monte-Carlo en Python. Les feuilles Scénarios et Sensibilité en reprennent les résultats sous forme de valeurs : elles ne se recalculent pas dans Excel. Pour les mettre à jour : uv run simulation/simulation.py puis uv run simulation/export_excel.py."),
        (F_TEXTE, "Hypothèse simplificatrice des feuilles en formules : le parc est à l'âge moyen de sa vie (pas de dynamique de cohortes) et le vol est assuré."),
        (F_TEXTE, "Tous les montants sont illustratifs, sauf les kilométrages, les prix Véligo et les ordres de grandeur du vol et des batteries, sourcés dans la feuille Hypothèses."),
    ]
    for i, (police, texte) in enumerate(lignes, start=2):
        c = ws.cell(i, 2, texte)
        c.font = police or F_TEXTE
        c.alignment = Alignment(wrap_text=True, vertical="top")


def feuille_hypotheses(wb):
    ws = wb.create_sheet("Hypothèses")
    ws["A1"].value, ws["A1"].font = "Hypothèses du modèle", F_TITRE
    ws["A2"].value, ws["A2"].font = "Modifier uniquement la colonne C (texte bleu). Les noms de la colonne F sont utilisés dans les formules des autres feuilles.", F_NOTE
    entetes(ws, 4, ["Catégorie", "Paramètre", "Valeur", "Unité", "Source ou commentaire", "Nom"], [14, 52, 12, 18, 70, 24])
    ws.freeze_panes = "A5"
    for i, (cat, nom, libelle, valeur, unite, fmt, cle, source) in enumerate(HYPOTHESES, start=5):
        ws.cell(i, 1, cat).font = F_TEXTE
        ws.cell(i, 2, libelle).font = F_TEXTE
        c = ws.cell(i, 3, valeur)
        c.font, c.number_format = F_SAISIE, fmt
        if cle:
            c.fill = FOND_CLE
        ws.cell(i, 4, unite).font = F_TEXTE
        s = ws.cell(i, 5, source)
        s.font, s.alignment = F_NOTE, Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 6, nom).font = F_NOTE
        for col in range(1, 7):
            ws.cell(i, col).border = BORDURE
        wb.defined_names[nom] = DefinedName(nom, attr_text=f"'Hypothèses'!$C${i}")


# Lignes de la feuille Cotisation : (libellé, formule léger, formule intensif, format, style)
# {s} est remplacé par leger ou intensif, et {C} par la colonne de la même ligne.
def lignes_cotisation():
    return [
        ("Intermédiaires", None, None, None, "section"),
        ("Occupation d'équilibre du parc", "=remplissage/(churn+remplissage-churn*remplissage)", None, PCT, ""),
        ("Kilomètres par vélo du parc", "={C}5*km_{s}", None, NB, ""),
        ("Âge moyen du parc", "=(vie_{s}-1)/2", None, DEC, ""),
        ("Heures d'atelier par vélo du parc", "=h_1000km_{s}*h_facteur/1000*{C}6*(1+vieillissement*{C}7)", None, DEC, ""),
        ("Coût annuel attendu par vélo du parc", None, None, None, "section"),
        ("Forfait de suivi", "=forfait_suivi*({C}5+forfait_vacant*(1-{C}5))", None, EUR, ""),
        ("Main d'œuvre", "={C}8*taux_horaire", None, EUR, ""),
        ("Prime de durabilité réparateur", "=prime_rep*({C}10+{C}11)*score_moyen/100", None, EUR, ""),
        ("Pièces d'usure", "=pieces_km_{s}*pieces_facteur*{C}6", None, EUR, ""),
        ("Batterie", "=0", "=batterie_remplacements/vie_intensif*batterie_facteur*batterie_cout", EUR, ""),
        ("Rente d'usage fabricant", "=IF(taux_financement>0,cout_{s}*taux_financement/(1-(1+taux_financement)^(-vie_{s})),cout_{s}/vie_{s})", None, EUR, ""),
        ("Prime de durabilité fabricant", "=prime_fab*{C}15*score_moyen/75", None, EUR, ""),
        ("Casse prématurée", "=casse_{s}*casse_facteur*{C}5*casse_cout", None, EUR, ""),
        ("Vol : assurance et franchise, moins franchise refacturée", "=assurance_taux*cout_{s}+vol_{s}*({C}5+0.3*(1-{C}5))*assurance_franchise*cout_{s}-vol_{s}*{C}5*franchise_client_{s}*recouvrement_franchise", None, EUR, ""),
        ("Capteur kilométrique", "=capteur_cout*(1/vie_leger+capteur_renouvellement)", "=0", EUR, ""),
        ("Remise en état entre deux locataires", "=churn*{C}5*recond_{s}", None, EUR, ""),
        ("Coût par vélo du parc", "=SUM({C}10:{C}20)", None, EUR, "total"),
        ("Tarif de location", None, None, None, "section"),
        ("Part du coût liée aux kilomètres", "={C}11+{C}13+prime_rep*{C}11*score_moyen/100", None, EUR, ""),
        ("Part du coût fixe (+ structure fixe répartie)", "={C}21-{C}23+gestion_fixe/(n_leger+n_intensif)", None, EUR, ""),
        ("Coefficient gestion + provision", "=1+gestion_taux+provision", None, DEC, ""),
        ("Part fixe du tarif", "={C}24*{C}25/({C}5*(1-impayes))", None, EUR, "total"),
        ("Part variable du tarif, par km déclaré", "={C}23*{C}25/({C}5*(1-impayes)*km_{s}*(1-fraude_part*fraude_sous_declaration))", None, EUR3, "total"),
        ("Cotisation annuelle au kilométrage moyen", "={C}26+{C}27*km_{s}", None, EUR, "total"),
        ("Cotisation mensuelle", "={C}28/12", None, EUR, "total"),
        ("Recettes attendues par vélo loué (impayés et fraude déduits)", "=({C}26+{C}27*km_{s}*(1-fraude_part*fraude_sous_declaration))*(1-impayes)", None, EUR, ""),
        ("Comparaison", None, None, None, "section"),
        ("Cotisation affichée sur la page initiale", "=page_{s}", None, EUR, "lien"),
        ("Écart avec la page initiale", "={C}28/{C}32-1", None, PCT, ""),
    ]


def feuille_cotisation(wb):
    ws = wb.create_sheet("Cotisation")
    ws["A1"].value, ws["A1"].font = "Coût par vélo et tarif de location", F_TITRE
    ws["A2"].value, ws["A2"].font = "Valeurs attendues en régime courant, calculées à partir de la feuille Hypothèses.", F_NOTE
    entetes(ws, 3, ["Poste", "", "Usage léger", "Usage intensif"], [58, 2, 18, 18])
    ws.freeze_panes = "A4"
    for i, (libelle, f_l, f_i, fmt, style) in enumerate(lignes_cotisation(), start=4):
        a = ws.cell(i, 1, libelle)
        if style == "section":
            for col in range(1, 5):
                ws.cell(i, col).fill = FOND_SECTION
            a.font = F_SOUS
            continue
        a.font = F_TOTAL if style == "total" else F_TEXTE
        for col, seg, f in (("C", "leger", f_l), ("D", "intensif", f_i if f_i else f_l)):
            c = ws[f"{col}{i}"]
            c.value = f.format(s=seg, C=col)
            c.number_format = fmt
            c.font = F_LIEN if style == "lien" else (F_TOTAL if style == "total" else F_TEXTE)
            if style == "total":
                c.fill = FOND_TOTAL
            c.border = BORDURE
    return ws


def feuille_fonds(wb):
    ws = wb.create_sheet("Fonds 10 ans")
    ws["A1"].value, ws["A1"].font = "Projection du fonds commun sur 10 ans (valeurs attendues)", F_TITRE
    ws["A2"].value, ws["A2"].font = ("Sans aléa : pour le risque, voir la feuille Scénarios. Le tarif de l'année N est révisé "
                                     "d'après le résultat de l'année N-1, dans la limite des plafonds, si la révision est activée."), F_NOTE
    ws.column_dimensions["A"].width = 46
    ws.cell(4, 1, "Année").font = F_ENTETE
    ws.cell(4, 1).fill = FOND_ENTETE
    for t in range(1, 11):
        col = get_column_letter(t + 1)
        ws.column_dimensions[col].width = 13
        c = ws.cell(4, t + 1, t)
        c.font, c.fill, c.number_format = F_ENTETE, FOND_ENTETE, "0"
        c.alignment = Alignment(horizontal="right")

    lignes = [
        ("Indice d'inflation des coûts", "=(1+inflation)^({col}4-1)", DEC, ""),
        ("Coût du parc léger", "=Cotisation!$C$21*n_leger*{col}5", EUR, "lien"),
        ("Coût du parc intensif", "=Cotisation!$D$21*n_intensif*{col}5", EUR, "lien"),
        ("Chargement de gestion (+ structure fixe)", "=gestion_taux*({col}6+{col}7)+gestion_fixe*{col}5", EUR, ""),
        ("Coûts totaux", "=SUM({col}6:{col}8)", EUR, "total"),
        ("Multiplicateur du tarif", None, DEC, ""),
        ("Recettes, parc léger", "=Cotisation!$C$30*n_leger*Cotisation!$C$5*{col}10", EUR, "lien"),
        ("Recettes, parc intensif", "=Cotisation!$D$30*n_intensif*Cotisation!$D$5*{col}10", EUR, "lien"),
        ("Recettes totales", "={col}11+{col}12", EUR, "total"),
        ("Résultat de l'année", "={col}13-{col}9", EUR, "total"),
        ("Marge", "=IF({col}13=0,0,{col}14/{col}13)", PCT, ""),
        ("Réserve en fin d'année", None, EUR, "total"),
        ("Cotisation mensuelle, usage léger", "=Cotisation!$C$29*{col}10", EUR, "lien"),
        ("Cotisation mensuelle, usage intensif", "=Cotisation!$D$29*{col}10", EUR, "lien"),
        ("Versements aux ateliers (forfait, main d'œuvre, prime), par atelier",
         "=(SUM(Cotisation!$C$10:$C$12)*n_leger+SUM(Cotisation!$D$10:$D$12)*n_intensif)*{col}5/n_ateliers", EUR, "lien"),
        ("Versements aux fabricants (rente et prime)",
         "=(Cotisation!$C$15+Cotisation!$C$16)*n_leger+(Cotisation!$D$15+Cotisation!$D$16)*n_intensif", EUR, "lien"),
    ]
    for i, (libelle, f, fmt, style) in enumerate(lignes, start=5):
        a = ws.cell(i, 1, libelle)
        a.font = F_TOTAL if style == "total" else F_TEXTE
        for t in range(1, 11):
            col = get_column_letter(t + 1)
            prec = get_column_letter(t)
            if libelle == "Multiplicateur du tarif":
                formule = "=1" if t == 1 else (f"=IF(revision=1,{prec}10*MIN(1+cap_hausse,MAX(1-cap_baisse,"
                                                f"{prec}9*(1+provision)/{prec}13)),{prec}10)")
            elif libelle == "Réserve en fin d'année":
                formule = f"=capital_initial+{col}14" if t == 1 else f"={prec}16+{col}14"
            else:
                formule = f.format(col=col)
            c = ws.cell(i, t + 1, formule)
            c.number_format = fmt
            c.font = F_LIEN if style == "lien" else (F_TOTAL if style == "total" else F_TEXTE)
            if style == "total":
                c.fill = FOND_TOTAL
            c.border = BORDURE
    ws.freeze_panes = "B5"
    n = len(lignes) + 6
    ws.cell(n, 1, "Réserve minimale sur 10 ans").font = F_TOTAL
    c = ws.cell(n, 2, "=MIN(B16:K16)")
    c.number_format, c.font = EUR, F_TOTAL
    ws.cell(n + 1, 1, "Marge cumulée sur 10 ans").font = F_TOTAL
    c = ws.cell(n + 1, 2, "=SUM(B14:K14)/SUM(B13:K13)")
    c.number_format, c.font = PCT, F_TOTAL


def feuille_scenarios(wb, res):
    ws = wb.create_sheet("Scénarios")
    ws["A1"].value, ws["A1"].font = f"{len(res['scenarios'])} scénarios critiques — résultats Monte-Carlo", F_TITRE
    ws["A2"].value = (f"Valeurs issues de simulation.py ({res['trajectoires_simulees']:,} trajectoires de 10 ans, "
                      f"graine {res['graine']}) : elles ne se recalculent pas dans Excel.").replace(",", " ")
    ws["A2"].font = F_NOTE
    colonnes = [
        ("Code", 7, None, "code"), ("Scénario", 40, None, "titre"), ("Description", 60, None, "description"),
        ("Ruine sur 10 ans", 11, PCT, "p_ruine"), ("Ruine sur 3 ans", 11, PCT, "p_ruine_3ans"),
        ("Marge moyenne", 11, PCT, "marge_moyenne"), ("Capital requis (99 %)", 15, EUR, "capital_requis_99"),
        ("Capital requis par vélo", 12, EUR, "capital_requis_par_velo"),
        ("Tarif d'équilibre (multiplicateur)", 13, DEC, "tarif_equilibre_mult"),
        ("Résultat annuel moyen", 14, EUR, "resultat_annuel_moyen"),
        ("Pire année (1 % des cas)", 14, EUR, "pire_annee_p01"),
        ("Réserve finale médiane", 15, EUR, "reserve_finale_p50"),
        ("Réserve finale (5 % des cas)", 15, EUR, "reserve_finale_p05"),
        ("Occupation moyenne", 11, PCT, "occupation_moyenne"),
        ("Léger an 1 (€/mois)", 11, EUR, "facture_leger_mois_an1"),
        ("Cargo an 1 (€/mois)", 11, EUR, "facture_intensif_mois_an1"),
        ("Léger an 10, médiane (€/mois)", 12, EUR, "facture_leger_mois_an10_p50"),
        ("Cargo an 10, médiane (€/mois)", 12, EUR, "facture_intensif_mois_an10_p50"),
        ("Cargo an 10, 95 % des cas (€/mois)", 12, EUR, "facture_intensif_mois_an10_p95"),
        ("Fonds collecté an 1", 14, EUR, "fonds_collecte_an1"),
        ("Revenu par atelier an 1", 13, EUR, "revenu_atelier_an1"),
    ]
    entetes(ws, 4, [c[0] for c in colonnes], [c[1] for c in colonnes])
    ws.row_dimensions[4].height = 45
    for i, sc in enumerate(res["scenarios"], start=5):
        for j, (_, _, fmt, cle) in enumerate(colonnes, start=1):
            v = sc.get(cle)
            c = ws.cell(i, j, "aucun" if (cle == "tarif_equilibre_mult" and v is None) else v)
            c.font = F_TEXTE
            if fmt:
                c.number_format = fmt
            if cle == "description":
                c.alignment = ENVELOPPE
            c.border = BORDURE
    fin = 4 + len(res["scenarios"])
    ws.conditional_formatting.add(f"D5:E{fin}", ColorScaleRule(start_type="num", start_value=0, start_color="E8F5E9",
                                                              mid_type="num", mid_value=0.1, mid_color="FFF3C4",
                                                              end_type="num", end_value=1, end_color="F4B6B0"))
    ws.cell(fin + 2, 1, "Tarif d'équilibre : plus petit multiplicateur du tarif de départ qui garde la ruine sous 1 % ; "
                        "« aucun » quand les usagers partent plus vite que le prix ne monte.").font = F_NOTE
    ws.cell(fin + 3, 1, "Ruine : la réserve devient négative au moins une fois, sans capital de départ. "
                        "Capital requis : capital de départ qui évite la ruine dans 99 % des trajectoires.").font = F_NOTE
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(colonnes))}{fin}"


LIBELLES_SENS = {
    "km_facteur": "Kilométrage (multiplicateur)", "churn": "Départs de locataires par an",
    "remplissage": "Relocation des vélos vacants", "impayes": "Impayés", "fraude_part": "Part de fraudeurs",
    "cout_leger": "Coût de fabrication, vélo classique (€)", "cout_intensif": "Coût de fabrication, vélo-cargo (€)",
    "taux_financement": "Taux de financement de la rente", "taux_horaire": "Taux horaire moyen (€)",
    "h_facteur": "Heures de main d'œuvre (multiplicateur)", "vieillissement": "Vieillissement (heures par année d'âge)",
    "pieces_facteur": "Consommation de pièces (multiplicateur)", "batterie_facteur": "Remplacement des batteries (multiplicateur)",
    "casse_facteur": "Casse prématurée (multiplicateur)", "serie_proba": "Défauts de série (par fabricant et par an)",
    "faillite_fab": "Faillite (par fabricant et par an)", "reforme_orphelins": "Réforme des vélos orphelins (par an)",
    "vol_leger": "Vol, vélo classique", "vol_intensif": "Vol, vélo-cargo", "inflation": "Inflation",
    "gestion_fixe": "Coûts de structure fixes (€/an)", "score_moyen": "Score moyen (primes)",
}


def feuille_sensibilite(wb, res):
    ws = wb.create_sheet("Sensibilité")
    sens = res["sensibilite"]
    ws["A1"].value, ws["A1"].font = "Analyse de sensibilité globale", F_TITRE
    ws["A2"].value = (f"{sens['n_jeux']:,} jeux de paramètres tirés en hypercube latin × {sens['n_traj_par_jeu']} trajectoires, "
                      "tarif de référence et révision annuelle. Valeurs issues de simulation.py.").replace(",", " ")
    ws["A2"].font = F_NOTE
    resume = [
        ("Probabilité de ruine moyenne sur l'espace exploré", sens["p_ruine_globale"], PCT),
        ("Jeux dont la ruine dépasse 5 %", sens["part_jeux_ruine_sup_5pct"], PCT),
        ("Jeux dont la ruine dépasse 50 %", sens["part_jeux_ruine_sup_50pct"], PCT),
        ("Marge médiane", sens["marge_mediane"], PCT),
        ("Capital requis (99 %), médiane des jeux", sens["capital_requis_99_median"], EUR),
        ("Capital requis (99 %), 90e centile des jeux", sens["capital_requis_99_p90"], EUR),
    ]
    for i, (lib, v, fmt) in enumerate(resume, start=4):
        ws.cell(i, 1, lib).font = F_TEXTE
        c = ws.cell(i, 2, v)
        c.font, c.number_format = F_TEXTE, fmt

    debut = 12
    entetes(ws, debut, ["Rang", "Paramètre", "Minimum exploré", "Maximum exploré", "Corrélation de rang avec la marge",
                        "Ruine, quintile bas", "Ruine, quintile haut"], [8, 46, 14, 14, 16, 14, 14])
    ws.row_dimensions[debut].height = 32
    for k, r in enumerate(sens["classement"], start=1):
        ligne = debut + k
        valeurs = [k, LIBELLES_SENS.get(r["parametre"], r["parametre"]), r["plage"][0], r["plage"][1],
                   r["spearman_marge"], r["p_ruine_quintile_bas"], r["p_ruine_quintile_haut"]]
        for j, v in enumerate(valeurs, start=1):
            c = ws.cell(ligne, j, v)
            c.font = F_TEXTE
            c.border = BORDURE
            c.number_format = {3: "#,##0.###", 4: "#,##0.###", 5: DEC, 6: PCT, 7: PCT}.get(j, "General")
    fin = debut + len(sens["classement"])
    ws.conditional_formatting.add(f"E{debut + 1}:E{fin}", ColorScaleRule(
        start_type="num", start_value=-0.6, start_color="F4B6B0", mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=0.6, end_color="B7DFC4"))

    g = sens["grille"]
    ligne = fin + 3
    ws.cell(ligne, 1, "Probabilité de ruine selon les deux paramètres les plus influents").font = F_SOUS
    ws.cell(ligne + 1, 1, f"Lignes : {LIBELLES_SENS[g['x']]} ; colonnes : {LIBELLES_SENS[g['y']]}").font = F_NOTE
    bx, by = g["bornes_x"], g["bornes_y"]
    ent = ligne + 2
    ws.cell(ent, 2, "").fill = FOND_ENTETE
    for j in range(4):
        c = ws.cell(ent, 3 + j, f"{by[j]:.2f} à {by[j + 1]:.2f}".replace(".", ","))
        c.font, c.fill = F_ENTETE, FOND_ENTETE
    for i in range(4):
        c = ws.cell(ent + 1 + i, 2, f"{bx[i]:,.0f} à {bx[i + 1]:,.0f}".replace(",", " "))
        c.font = F_TOTAL
        for j in range(4):
            v = ws.cell(ent + 1 + i, 3 + j, g["p_ruine"][i][j])
            v.font, v.number_format = F_TEXTE, PCT
    ws.conditional_formatting.add(f"C{ent + 1}:F{ent + 4}", ColorScaleRule(
        start_type="num", start_value=0.4, start_color="FFF3C4", end_type="num", end_value=1, end_color="F4B6B0"))


def feuille_vacance(wb, res):
    ws = wb.create_sheet("Vacance prudente")
    ws["A1"].value, ws["A1"].font = "Tarifer avec une vacance prudente", F_TITRE
    ws["A2"].value = ("Le tarif de départ est calculé en supposant un taux d'occupation plus bas que l'équilibre (94 %). "
                      "La réalité est ensuite tirée au sort selon cinq situations. Valeurs issues de simulation.py.")
    ws["A2"].font = F_NOTE
    lignes = res["vacance_prudente"]
    realites = list(lignes[0]["realites"])
    ligne = 4
    for elasticite, titre in ((1.0, "Usagers sensibles au prix (élasticité 1)"), (0.0, "Usagers captifs (élasticité 0)")):
        ws.cell(ligne, 1, titre).font = F_SOUS
        ligne += 1
        entetes(ws, ligne, ["Occupation supposée", "Léger (€/mois)", "Cargo (€/mois)"]
                + [f"Ruine — {r}" for r in realites] + [f"Occupation réelle — {r}" for r in realites],
                [22, 12, 12] + [15] * (2 * len(realites)))
        ws.row_dimensions[ligne].height = 45
        debut = ligne + 1
        for l in (x for x in lignes if x["elasticite"] == elasticite):
            ligne += 1
            valeurs = [l["occupation_tarif"], l["leger_mois"], l["intensif_mois"]]
            valeurs += [l["realites"][r]["p_ruine"] for r in realites]
            valeurs += [l["realites"][r]["occupation"] for r in realites]
            for j, v in enumerate(valeurs, start=1):
                c = ws.cell(ligne, j, v)
                c.font, c.border = F_TEXTE, BORDURE
                c.number_format = "General" if j == 1 else (EUR if j <= 3 else PCT)
        fin_ruine = get_column_letter(3 + len(realites))
        ws.conditional_formatting.add(f"D{debut}:{fin_ruine}{ligne}", ColorScaleRule(
            start_type="num", start_value=0, start_color="E8F5E9", mid_type="num", mid_value=0.1, mid_color="FFF3C4",
            end_type="num", end_value=1, end_color="F4B6B0"))
        ligne += 2
    ws.cell(ligne, 1, "Lecture : supposer 80 % d'occupation au lieu de 94 % relève les prix d'environ 5 % et réduit fortement "
                      "le risque dans les tempêtes ; la révision annuelle fait ensuite redescendre le tarif si la vacance "
                      "ne se produit pas.").font = F_NOTE


def main():
    res = json.loads((ICI / "resultats" / "resultats.json").read_text())
    wb = Workbook()
    feuille_lisez_moi(wb, res)
    feuille_hypotheses(wb)
    feuille_cotisation(wb)
    feuille_fonds(wb)
    feuille_scenarios(wb, res)
    feuille_sensibilite(wb, res)
    feuille_vacance(wb, res)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb["Hypothèses"]["C16"].comment = Comment(
        "Paramètre le plus influent de l'analyse de sensibilité. Prix publics des vélos-cargos électriques "
        "français : 3 500 à 7 000 €. Scénarios S09 et S10 : 2 800 €.", "Pignon commun")
    sortie = ICI / "simulation-economique.xlsx"
    wb.save(sortie)
    print(sortie)


if __name__ == "__main__":
    main()
